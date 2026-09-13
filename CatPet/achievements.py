# -*- coding: utf-8 -*-
"""成就目录、解锁判定与展示控件。"""
import os
import tkinter as tk
from tkinter import ttk


DEFAULT_ACHIEVEMENTS = {
    'dummy': {
        'id': 'dummy', 'name': '笨蛋', 'desc': '开局就获得的成就，笨蛋也有好运气。',
        'condition': 'gift', 'value': 0, 'text_color': '#555555',
        'outline_color': '#b0b0b0', 'gift': True,
    },
    'rich_cat': {
        'id': 'rich_cat', 'name': '小富翁', 'desc': '总计财富达到 2000 G 后解锁。',
        'condition': 'wealth', 'value': 2000, 'text_color': '#8B6914',
        'outline_color': '#3b2a00', 'gift': False,
    },
}


class AchievementService:
    @staticmethod
    def _as_bool(value, default=False):
        if value in (None, ''):
            return default
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in ('1', 'true', 'yes', 'y', '是', '可', '可以')

    @classmethod
    def load_catalog(cls, asset_dir):
        catalog = {key: dict(value) for key, value in DEFAULT_ACHIEVEMENTS.items()}
        path = os.path.join(asset_dir, '成就簿.xlsx')
        if not os.path.exists(path):
            return catalog
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
            sheet = workbook.active
            headers = [str(cell.value).strip() if cell.value is not None else ''
                       for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            columns = {name: index for index, name in enumerate(headers) if name}
            required = ('成就ID', '名称', '描述', '解锁条件', '条件值',
                        '文字颜色', '描边颜色', '开局赠送')
            result = {}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                def cell(name, default=''):
                    index = columns.get(name)
                    if index is None or index >= len(row) or row[index] is None:
                        return default
                    return row[index]
                achievement_id = str(cell('成就ID')).strip()
                if not achievement_id:
                    continue
                try:
                    value = int(float(cell('条件值', 0) or 0))
                except (TypeError, ValueError):
                    value = 0
                result[achievement_id] = {
                    'id': achievement_id,
                    'name': str(cell('名称', achievement_id)).strip(),
                    'desc': str(cell('描述', '')).strip(),
                    'condition': str(cell('解锁条件', 'gift')).strip().lower(),
                    'comparison': str(cell('解锁比较', '>=')).strip() or '>=',
                    'value': value,
                    'text_color': str(cell('文字颜色', '#ffffff')).strip(),
                    'outline_color': str(cell('描边颜色', '#333333')).strip(),
                    'gift': cls._as_bool(cell('开局赠送'), False),
                }
            if result and all(name in columns for name in required):
                catalog = result
            workbook.close()
        except Exception:
            pass
        return catalog

    @staticmethod
    def total_wealth(state, items):
        total = max(0, int(getattr(state, 'gold', 0)))
        total += sum(max(0, int(t.get('value', 0)))
                     for t in getattr(state, 'treasures', []))
        for item_id, count in getattr(state, 'inventory', {}).items():
            count = max(0, int(count))
            if not count:
                continue
            info = items.get(item_id) or {}
            value = info.get('sell_price')
            if value is None:
                value = info.get('buy_price', info.get('price', 0))
            total += max(0, int(value or 0)) * count
        return total

    @staticmethod
    def _compare(actual, comparison, target):
        actual = float(actual)
        target = float(target)
        if comparison in ('>', 'gt'):
            return actual > target
        if comparison in ('<', 'lt'):
            return actual < target
        if comparison in ('<=', 'le'):
            return actual <= target
        if comparison in ('==', 'eq'):
            return actual == target
        return actual >= target

    @classmethod
    def _metric(cls, state, items, condition):
        stats = getattr(state, 'stats', {}) or {}
        inventory = getattr(state, 'inventory', {}) or {}
        equipment_count = sum(
            1 for item_id, count in inventory.items()
            if int(count or 0) > 0
            and (items.get(item_id) or {}).get('category') in
            ('头饰', '服装', '饰品'))
        purchased = stats.get('purchased_items', {})
        fashion_items = int(all(
            int(purchased.get(item_id, 0) or 0) > 0
            for item_id in ('hairstyle_tool', 'hair_dye', 'lens')))
        summary = {
            'wealth': cls.total_wealth(state, items),
            'gold_zero_seen': int(stats.get('gold_zero_seen', 0)),
            'emotion_max': int(stats.get('emotion_max', 0)),
            'emotion_min': int(stats.get('emotion_min', 100)),
            'emotion_exact10': int(stats.get('emotion_exact10', 0)),
            'hp_min': int(stats.get('hp_min', 0)),
            'equipment_count': equipment_count,
            'jump_guide': int(inventory.get('jump_guide', 0)),
            'pat_guide': int(inventory.get('pat_guide', 0)),
            'adventure_early_fail_count': int(stats.get('adventure_early_fail_count', 0)),
            'strength': int(getattr(state, 'strength', 0)),
            'wisdom': int(getattr(state, 'wisdom', 0)),
            'agility': int(getattr(state, 'agility', 0)),
            'max_hp': int(getattr(state, 'max_hp', 0)),
            'cat_click_count': int(stats.get('cat_click_count', 0)),
            'purchase_count': int(stats.get('purchase_count', 0)),
            'treasure_sold_count': int(stats.get('treasure_sold_count', 0)),
            'online_minutes': int(stats.get('online_seconds', 0) // 60),
            'level': int(getattr(state, 'level', 1)),
            'fashion_items': fashion_items,
            'adventure_count': int(stats.get('adventure_count', 0)),
            'forest_adventure_count': int(stats.get('forest_adventure_count', 0)),
            'adventure_minutes': int(stats.get('adventure_minutes', 0)),
            'max_feed_amount': int(stats.get('max_feed_amount', 0)),
            'mushroom_damage_count': int(stats.get('mushroom_damage_count', 0)),
        }
        return summary.get(condition, 0)

    @classmethod
    def sync_unlocks(cls, state, catalog, items):
        if not isinstance(getattr(state, 'achievements', None), set):
            state.achievements = set(getattr(state, 'achievements', []) or [])
        newly = set()
        for achievement_id, info in catalog.items():
            if info.get('gift'):
                if achievement_id not in state.achievements:
                    state.achievements.add(achievement_id)
                    newly.add(achievement_id)
                continue
            if achievement_id in state.achievements:
                continue
            condition = str(info.get('condition', '')).strip().lower()
            target = int(info.get('value', 0))
            comparison = str(info.get('comparison', '>=')).strip()
            actual = cls._metric(state, items, condition)
            if condition and cls._compare(actual, comparison, target):
                state.achievements.add(achievement_id)
                newly.add(achievement_id)
        unlocked = [key for key in catalog if key in state.achievements]
        selected = getattr(state, 'selected_achievement', None)
        if selected not in unlocked:
            state.selected_achievement = unlocked[0] if unlocked else None
        return newly


EXPRESSION_BOOK_UNLOCK_ACHIEVEMENTS = {
    'expression_secret': 'pain_cat',
    'expression_comfy': 'bliss_cat',
    'expression_grumpy': 'impatient_cat',
}


class AchievementMixin:
    def _sync_expression_book_unlocks(self):
        changed = False
        unlocked = getattr(self.game, 'achievements', set())
        inventory = getattr(self.game, 'inventory', {})
        for book_id, achievement_id in EXPRESSION_BOOK_UNLOCK_ACHIEVEMENTS.items():
            if achievement_id in unlocked and int(inventory.get(book_id, 0)) <= 0:
                inventory[book_id] = 1
                changed = True
        return changed

    def _build_achievement_selector(self, parent, bg):
        frame = tk.Frame(parent, bg=bg)
        frame.pack(fill='x', pady=(0, 4))
        row = tk.Frame(frame, bg=bg)
        row.pack(anchor='center')
        self._achievement_canvas = tk.Canvas(
            row, width=250, height=34, bg=bg, highlightthickness=0,
            cursor='hand2')
        self._achievement_canvas.pack(anchor='center')
        self._achievement_canvas.bind(
            '<Button-1>', self._open_achievement_menu)
        self._achievement_canvas.bind(
            '<Enter>', self._show_achievement_tooltip)
        self._achievement_canvas.bind(
            '<Leave>', self._hide_achievement_tooltip)
        self._achievement_tooltip = None
        self._achievement_names = {}
        self._refresh_achievement_ui()

    def _current_achievement(self):
        catalog = getattr(self, 'achievement_catalog', {})
        selected = getattr(self.game, 'selected_achievement', None)
        return catalog.get(selected)

    def _refresh_achievement_ui(self):
        canvas = getattr(self, '_achievement_canvas', None)
        if canvas is None:
            return
        try:
            if not canvas.winfo_exists():
                return
        except Exception:
            return
        catalog = getattr(self, 'achievement_catalog', {})
        newly = AchievementService.sync_unlocks(
            self.game, catalog, globals().get('ITEMS', {}))
        unlocked = [key for key in catalog if key in self.game.achievements]
        self._achievement_names = {key: catalog[key] for key in unlocked}
        current = self._current_achievement()
        canvas.delete('all')
        width = max(250, canvas.winfo_width())
        font = ('Microsoft YaHei UI', 14, 'bold')
        quote_font = ('Microsoft YaHei UI', 16, 'bold')

        def draw_quoted(text, fill, outline=None):
            if outline is not None:
                for dx, dy in ((-1, 0), (1, 0), (0, -1), (0, 1),
                               (-1, -1), (-1, 1), (1, -1), (1, 1)):
                    canvas.create_text(width // 2 + dx, 17 + dy, text=text,
                                       fill=outline, font=font, anchor='center')
            text_id = canvas.create_text(
                width // 2, 17, text=text, fill=fill, font=font,
                anchor='center')
            left, _, right, _ = canvas.bbox(text_id)
            canvas.create_text(left - 4, 17, text='\u201c', fill='#000000',
                               font=quote_font, anchor='e')
            canvas.create_text(right + 4, 17, text='\u201d', fill='#000000',
                               font=quote_font, anchor='w')

        if current is None:
            draw_quoted('\u6682\u65e0\u6210\u5c31', '#888888')
        else:
            draw_quoted(current['name'], current['text_color'],
                        current['outline_color'])
        if newly:
            try:
                self._save_satiety()
            except Exception:
                pass

    def _open_achievement_menu(self, event=None):
        outer = getattr(self, '_adjust_wrap', None)
        try:
            page_available = (
                outer is not None and outer.winfo_exists()
                and getattr(self, '_warehouse_win', None) is not None
                and self._warehouse_win.winfo_exists())
        except Exception:
            page_available = False
        if page_available:
            self._show_achievement_selector_page()
            return
        self._open_achievement_popup_menu(event)


    def _open_achievement_popup_menu(self, event=None):
        canvas = getattr(self, '_achievement_canvas', None)
        if canvas is None:
            return
        menu = tk.Menu(canvas, tearoff=0, font=('Microsoft YaHei UI', 9))
        unlocked = [key for key in getattr(self, 'achievement_catalog', {})
                    if key in self.game.achievements]
        if not unlocked:
            menu.add_command(label='暂无成就', state='disabled')
        else:
            selected = getattr(self.game, 'selected_achievement', None)
            for achievement_id in unlocked:
                info = self.achievement_catalog[achievement_id]
                mark = ' ' if achievement_id == selected else ' '
                menu.add_command(
                    label=mark + info['name'],
                    command=lambda k=achievement_id: self._select_achievement(k))
        x = event.x_root if event is not None else canvas.winfo_rootx()
        y = (event.y_root if event is not None
             else canvas.winfo_rooty() + canvas.winfo_height())
        try:
            menu.tk_popup(x, y)
        finally:
            menu.grab_release()


    def _show_achievement_selector_page(self):
        outer = getattr(self, '_adjust_wrap', None)
        content = getattr(self, '_adjust_content', None)
        list_wrap = getattr(self, '_warehouse_list_wrap', None)
        subbar = getattr(self, '_equip_subbar', None)
        if outer is None or content is None:
            return
        self._hide_achievement_tooltip()
        for child in content.winfo_children():
            child.destroy()
        if subbar is not None:
            subbar.pack_forget()
        if list_wrap is not None:
            list_wrap.pack_forget()
        outer.pack(fill='both', expand=True, padx=14, pady=14)

        header = tk.Frame(content, bg=THEME['card'])
        header.pack(fill='x', pady=(0, 10))
        make_button(header, ' 返回', self._return_to_equipment_list,
                    kind='ghost', width=8).pack(side='left')
        make_label(header, '成就选择', style='title',
                   bg=THEME['card']).pack(side='left', padx=(12, 0))

        unlocked = [key for key in getattr(self, 'achievement_catalog', {})
                    if key in self.game.achievements]
        list_container = tk.Frame(content, bg=THEME['card'])
        list_container.pack(fill='both', expand=True)
        self._achievement_selector_list = list_container
        if not unlocked:
            make_label(list_container, '暂无已解锁成就', bg=THEME['card'],
                       fg=THEME['muted']).pack(anchor='w', pady=8)
            self._show_achievement_detail_placeholder()
            return

        selected = getattr(self.game, 'selected_achievement', None)
        for achievement_id in unlocked:
            info = self.achievement_catalog[achievement_id]
            row = make_card(list_container, bg=THEME['card'])
            row.pack(fill='x', pady=2)
            row._achievement_id = achievement_id
            row._quality_color = info.get('text_color', THEME['text'])
            row._quality_border_thick = False
            row._selected = achievement_id == selected
            mark = '☑ ' if achievement_id == selected else '☐ '
            label = make_label(
                row, text=mark + info['name'], bg=THEME['card'],
                fg=info.get('text_color', THEME['text']),
                font=('Microsoft YaHei UI', 11, 'bold'),
                cursor='hand2', anchor='w')
            label.pack(side='left', fill='x', expand=True,
                       padx=6, pady=4)
            row._achievement_label = label
            for widget in (row, label):
                widget.bind(
                    '<Button-1>',
                    lambda e, k=achievement_id:
                    self._choose_achievement_from_page(k))
            self._bind_achievement_detail_frame(row, info)
        layout = getattr(self, '_layout_two_columns', None)
        if callable(layout):
            layout(list_container)
        current = self._current_achievement()
        if current is None:
            current = self.achievement_catalog[unlocked[0]]
        self._show_achievement_detail(current, pinned=True)


    def _choose_achievement_from_page(self, achievement_id):
        if achievement_id not in self.game.achievements:
            return
        self.game.selected_achievement = achievement_id
        self._refresh_achievement_ui()
        try:
            self._save_satiety()
        except Exception:
            pass
        content = getattr(self, '_achievement_selector_list', None)
        if content is not None:
            for row in content.winfo_children():
                row_id = getattr(row, '_achievement_id', None)
                if row_id is None:
                    continue
                selected = row_id == achievement_id
                row._selected = selected
                label = getattr(row, '_achievement_label', None)
                if label is not None:
                    info = self.achievement_catalog[row_id]
                    mark = '☑ ' if selected else '☐ '
                    label.config(text=mark + info['name'])
            layout = getattr(self, '_layout_two_columns', None)
            if callable(layout):
                layout(content)
        info = self.achievement_catalog.get(achievement_id)
        if info is not None:
            self._show_achievement_detail(info, pinned=True)


    def _bind_achievement_detail_frame(self, widget, info):
        def show_detail(event=None):
            self._show_achievement_detail(info, pinned=True)

        def bind_tree(item):
            item.bind('<Enter>', show_detail, add='+')
            for child in item.winfo_children():
                bind_tree(child)

        bind_tree(widget)


    def _achievement_condition_text(self, info):
        condition = info.get('condition', '')
        comparison = info.get('comparison', '>=')
        value = info.get('value', 0)
        if condition == 'gift':
            return '开局赠送'
        if condition == 'emotion_exact10':
            return '情绪等于 10'
        labels = {
            'wealth': '总财富',
            'emotion_min': '情绪',
            'emotion_max': '情绪',
            'equipment_count': '拥有的装备数量',
            'jump_guide': '已获得猫咪跳远指南',
            'pat_guide': '已获得乖，摸摸头',
            'wisdom': '智慧',
            'strength': '力量',
            'agility': '敏捷',
            'max_hp': '最大活力',
            'cat_click_count': '点击猫猫次数',
            'purchase_count': '购买次数',
            'treasure_sold_count': '出售宝藏次数',
            'fashion_items': '养成道具购买进度',
            'online_minutes': '累计在线分钟数',
            'level': '等级',
            'adventure_early_fail_count': '提前失败次数',
            'adventure_count': '探险次数',
            'forest_adventure_count': '低语森林成功次数',
            'adventure_minutes': '累计探险分钟数',
            'max_feed_amount': '单次最大喂食数量',
            'mushroom_damage_count': '蘑菇吃坏肚子次数',
        }
        label = labels.get(condition, condition or '未知条件')
        return f'{label} {comparison} {value}'


    def _show_achievement_detail_placeholder(self):
        content = getattr(self, '_warehouse_detail_content', None)
        if content is None or not content.winfo_exists():
            return
        for child in content.winfo_children():
            child.destroy()
        card = getattr(self, '_warehouse_detail_card', None)
        if card is not None:
            card.configure(
                highlightbackground=THEME['border'],
                highlightcolor=THEME['border'], highlightthickness=1)
        make_label(
            content, '成就详情', bg=THEME['card'], fg=THEME['muted'],
            font=(FONT_FAMILY, 10, 'bold')).pack(
                anchor='w', padx=12, pady=(14, 6))
        make_label(
            content, '暂无已解锁成就', bg=THEME['card'],
            fg=THEME['muted'], font=(FONT_FAMILY, 9)).pack(
                anchor='w', padx=12)


    def _show_achievement_detail(self, info, pinned=False):
        if pinned:
            self._pinned_warehouse_detail = info
        job = getattr(self, '_warehouse_detail_job', None)
        if job is not None:
            try:
                self.root.after_cancel(job)
            except Exception:
                pass
        self._warehouse_detail_job = None
        self._cancel_warehouse_detail_hide()
        content = getattr(self, '_warehouse_detail_content', None)
        if content is None or not content.winfo_exists():
            return
        for child in content.winfo_children():
            child.destroy()
        text_color = info.get('text_color', THEME['text'])
        outline_color = info.get('outline_color', THEME['border'])
        card = getattr(self, '_warehouse_detail_card', None)
        if card is not None:
            card.configure(
                highlightbackground=outline_color,
                highlightcolor=outline_color, highlightthickness=2)
        make_label(
            content, '成就详情', bg=THEME['card'], fg=THEME['muted'],
            font=(FONT_FAMILY, 10, 'bold')).pack(
                anchor='w', padx=12, pady=(14, 5))
        make_label(
            content, info.get('name', '-'), bg=THEME['card'],
            fg=text_color, font=(FONT_FAMILY, 12, 'bold'),
            anchor='w', justify='left', wraplength=240).pack(
                fill='x', padx=12, pady=(0, 2))
        selected = getattr(self.game, 'selected_achievement', None)
        status = '当前选择' if selected == info.get('id') else '已解锁'
        make_label(
            content, f'状态：{status}', bg=THEME['card'],
            fg=THEME['muted'], font=(FONT_FAMILY, 9),
            anchor='w').pack(fill='x', padx=12)
        tk.Frame(content, bg=THEME['border'], height=1).pack(
            fill='x', padx=12, pady=(7, 6))
        make_label(
            content, info.get('desc', '（暂无描述）'), bg=THEME['card'],
            fg=THEME['text'], font=(FONT_FAMILY, 9), justify='left',
            anchor='w', wraplength=240).pack(fill='x', padx=12)
        make_label(
            content, '解锁条件', bg=THEME['card'], fg=THEME['muted'],
            font=(FONT_FAMILY, 9, 'bold'), anchor='w').pack(
                fill='x', padx=12, pady=(12, 2))
        make_label(
            content, self._achievement_condition_text(info),
            bg=THEME['card'], fg=THEME['text'],
            font=(FONT_FAMILY, 9), justify='left', anchor='w',
            wraplength=240).pack(fill='x', padx=12)

    def _select_achievement(self, achievement_id):
        if achievement_id not in self.game.achievements:
            return
        self.game.selected_achievement = achievement_id
        self._refresh_achievement_ui()
        try:
            self._save_satiety()
        except Exception:
            pass

    def _show_achievement_tooltip(self, event=None):
        info = self._current_achievement()
        if info is None:
            return
        tooltip = getattr(self, '_achievement_tooltip', None)
        if tooltip is None or not tooltip.winfo_exists():
            tooltip = tk.Toplevel(self.root)
            tooltip.overrideredirect(True)
            tooltip.attributes('-topmost', True)
            label = tk.Label(tooltip, text='', justify='left', relief='solid',
                             bd=1, bg='#fffbe8', fg='#333333', padx=8, pady=5,
                             font=('Microsoft YaHei UI', 9), wraplength=260)
            label.pack()
            self._achievement_tooltip = tooltip
            self._achievement_tooltip_label = label
        self._achievement_tooltip_label.config(text=info.get('desc', ''))
        widget = getattr(self, '_achievement_canvas', None)
        if widget is not None:
            x = widget.winfo_rootx()
            y = widget.winfo_rooty() + widget.winfo_height() + 4
            tooltip.geometry(f'+{x}+{y}')
        tooltip.deiconify()
        tooltip.lift()

    def _hide_achievement_tooltip(self, event=None):
        tooltip = getattr(self, '_achievement_tooltip', None)
        if tooltip is not None:
            try:
                tooltip.withdraw()
            except Exception:
                pass

    def _destroy_achievement_selector(self):
        self._hide_achievement_tooltip()
        tooltip = getattr(self, '_achievement_tooltip', None)
        if tooltip is not None:
            try:
                tooltip.destroy()
            except Exception:
                pass
        self._achievement_tooltip = None
        self._achievement_tooltip_label = None
        self._achievement_canvas = None
        self._achievement_names = {}
        self._achievement_selector_list = None

def install_runtime_globals(namespace):
    globals().update(namespace)
